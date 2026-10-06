from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_paragraphs(doc, text_list):
    for text in text_list:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.name = 'Calibri'


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')


def add_section(doc, title, content, bullet_items=None):
    doc.add_heading(title, level=1)
    add_paragraphs(doc, content)
    if bullet_items:
        add_bullets(doc, bullet_items)
    doc.add_page_break()


def build_report():
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    for paragraph in doc.paragraphs:
        paragraph.style = None

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run('SmartCampus – AI College Information Chatbot')
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.name = 'Calibri'

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run('Project Report')
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.name = 'Calibri'

    details = doc.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details_run = details.add_run('Department of Computer Science and Engineering\nAcademic Project Report\n2026')
    details_run.font.size = Pt(12)
    details_run.font.name = 'Calibri'
    doc.add_page_break()

    add_section(
        doc,
        'Abstract',
        [
            'SmartCampus is an intelligent college information system designed to make campus-related information easily accessible to students and administrators through a conversational interface. The project combines a responsive web application with a retrieval-based chatbot that answers natural-language questions about courses, fees, departments, faculty, announcements, scholarships, admissions, and academic procedures. The system is built using Flask for backend services, SQLite for database storage, SQLAlchemy for ORM-based persistence, HTML/CSS/JavaScript for the frontend, and scikit-learn for TF-IDF and cosine similarity matching. The aim is to provide a practical digital campus assistant that reduces the need for students to search through multiple websites or manual notices while ensuring that information remains easy to maintain through a secure admin interface.',
            'Unlike simple rule-based assistants, SmartCampus demonstrates real NLP behavior using a vector-space matching model. The chatbot evaluates the student query against a knowledge base of curated FAQs and matches the most relevant answer using the similarity score between the query vector and the FAQ vectors. If the similarity does not pass a confidence threshold, the system returns a fallback response instead of making a misleading claim. This makes the solution more reliable, explainable, and suitable for academic demonstration. The project also includes a student dashboard, admin analytics panel, chat history, and CRUD operations for core campus information records.',
            'The completed system addresses the need for student support in a digital campus environment. It enhances communication efficiency, reduces the burden on administrative staff, and provides a centralized place for institutional information. The report explains the project lifecycle from requirement analysis through design, implementation, testing, and future enhancements, highlighting how a real-world educational problem was transformed into a working web application with practical AI capabilities.'
        ]
    )

    add_section(
        doc,
        '1. Introduction',
        [
            'Higher education institutions generate and manage a large amount of information that students must access frequently. Admission criteria, scholarship opportunities, departmental contact points, fee schedules, announcements, exam schedules, faculty specialization, and course offerings are all examples of information that are vital for student success. In many institutions, this information is scattered across notice boards, PDF documents, websites, and manual helpdesk interactions. Students often waste valuable time navigating multiple sources. As educational environments become more digital and data-driven, there is a strong need for a centralized system that can provide accurate, direct, and well-organized campus information.',
            'SmartCampus addresses this challenge by acting as an AI-powered college information portal. The application enables students to ask questions in natural language and receive relevant answers in a conversational manner. It also allows administrators to manage the information repository without needing to edit raw files or write database queries manually. The combination of a web interface, secure authentication, structured college data, and a retrieval-based NLP engine makes the project useful, scalable, and academically relevant. The system is designed as a complete full-stack web app and serves as a practical example of how AI-based retrieval systems can be integrated into educational software.',
            'The project is inspired by common student needs in college life. Questions such as “What are the admission criteria?” “Which engineering courses do you offer?” and “How can I contact the department?” are highly representative of real user queries. By enabling the system to interpret these as meaningful requests and retrieve answers from a curated knowledge base, the project demonstrates how artificial intelligence can be applied to benefit both students and administrators. The educational value lies in the fusion of web development, database design, authentication, and information retrieval.'
        ]
    )

    add_section(
        doc,
        '2. Problem Statement',
        [
            'Students frequently require up-to-date and accurate information regarding academic programs, campus services, fee structures, and student support. However, this information is often fragmented, outdated, or not accessible through a single interface. A student may need to contact a department office for a simple answer or manually browse through multiple pages to locate a fee rule or scholarship requirement. This process is inefficient and can lead to confusion or missed deadlines. Additionally, many institutions do not offer any personalized digital assistance that can answer frequently asked questions in real time.',
            'Another issue is the inability of many static portals to understand user queries in natural language. A student may ask a question using synonyms or informal wording, and a traditional database search may not return the intended result. The system needs to support semantic flexibility, not just exact keyword matching. Moreover, when an answer cannot be confidently found, the system should avoid fabricating inaccurate information and instead provide a fallback response. This challenge requires both good data organization and a retrieval model that understands relevance rather than relying on simplistic if-else responses.',
            'The project problem statement therefore focuses on developing an AI college chatbot that offers structured information, supports student queries in natural language, handles authentication safely, and empowers administrators to maintain accurate knowledge content. The solution must be modular, maintainable, and fit within an academic environment where the goal is to build a demonstrable, fully functional system using common open-source technologies.'
        ]
    )

    add_section(
        doc,
        '3. Objectives of the Project',
        [
            'The primary objective of SmartCampus is to create a complete college information system that supports students with real-time, natural-language access to institutional information. The system should reduce the effort required to locate frequently asked questions and provide a more engaging campus experience. It should also simplify content management for administrators, allowing them to update knowledge records without technical complexity. The project aims to bring together frontend development, backend APIs, database management, security practices, and NLP into a single educational web application.',
            'The specific objectives are as follows: to build a responsive student portal with login and chat features; to implement a secure administrator dashboard for content management and analytics; to store campus information in a structured SQLite database; to integrate a retrieval-based NLP engine for matching student questions with the most relevant FAQ; to provide fallback responses when no confident answer exists; and to validate the system through functional and relevance-based testing. The project also focuses on modularity and clean coding practices so that future extensions such as voice support, multilingual handling, or integration with a real LLM can be added later.'
        ],
        [
            'Develop a robust student login and profile system.',
            'Provide natural-language campus information support using TF-IDF and cosine similarity.',
            'Create an admin module for managing FAQs, courses, departments, fees, and announcements.',
            'Maintain student chat history and provide analytics for admin review.',
            'Ensure secure authentication and fallback behavior for uncertain answers.'
        ]
    )

    add_section(
        doc,
        '4. Scope and Functional Requirements',
        [
            'The scope of the project includes both end-user functionality and administrative features. Students can register, log in, ask questions, view earlier chat conversations, and browse campus information such as departments, fee structures, faculty profiles, and announcements. The application allows students to receive answers through a conversational interface that resembles a support bot but operates using a real knowledge base and NLP-based retrieval system. The front-end is designed to be clean, responsive, and easy to navigate on desktop and mobile devices.',
            'The administrator module handles the management of knowledge content. Administrators can create, edit, and delete FAQs, academic programs, departments, faculty, fee schedules, and event announcements. The admin dashboard also provides summary metrics such as total students, total questions asked, total FAQ entries, category-wise question distribution, and recent chat activity. This ensures that the system is not only a chatbot but also a management tool for college information. The project scope also includes secure session handling and session-based authorization to ensure that only authenticated students and admins can access protected routes.',
            'Functional requirements are grouped into student-side features, admin-side features, and system features. Student-side features include authentication, chat with the bot, answer retrieval, and browsing campus-related information. Admin-side features include CRUD management and insights into student interaction data. System features include modular architecture, SQLite persistence, multilingual-ready text normalization, and fallback responses. Non-functional requirements emphasize usability, maintainability, data integrity, performance, and security of password handling.'
        ]
    )

    add_section(
        doc,
        '5. Requirements Analysis',
        [
            'The requirement analysis was performed with a focus on the real needs of a college environment and the practical constraints of a student project. It was important to ensure that the application would support authentic information retrieval rather than static canned answers. Therefore, the system was designed around an editable FAQ corpus that acts as the source of truth for the bot. Administrators can refine the content whenever policies, fee structures, or programs change. The parser and retrieval model were designed to handle majority of common student questions without requiring a large, expensive language model or external API.',
            'The project follows a balanced architecture: a lightweight but capable backend, a simple SQLite database, and a standards-based web frontend. This approach is suitable for educational demonstration because it reduces setup complexity while still enabling real functionality. The requirement analysis also emphasized secure handling of user passwords, storing hashes rather than plaintext values, and maintaining separate roles for students and administrators. This is essential because unauthorized access to campus data or admin controls would undermine trust in the system.',
            'From the users’ perspective, the application must feel responsive and dependable. The chatbot should return answers quickly and consistently, even when questions are phrased differently than the canonical FAQ wording. The database must preserve student chat history and allow administrators to review insights. The design also ensures that the system has a graceful fallback when the model cannot determine a confident answer. This requirement is especially important for educational chatbot deployment because inaccurate answers can reduce confidence in the system.'
        ],
        [
            'The system must support student login and registration.',
            'The chatbot must use natural-language input rather than only prewritten menu options.',
            'The admin panel must support create, read, update, and delete operations.',
            'The system must store chat data per student and preserve historical context.',
            'When confidence is low, the chatbot must display a safe fallback response.'
        ]
    )

    add_section(
        doc,
        '6. System Design and Architecture',
        [
            'SmartCampus follows a modular three-layer architecture. The presentation layer is composed of the HTML, CSS, and JavaScript files that provide the user interface for student and admin workflows. The application shell loads the dashboard, login forms, chatbot panel, and admin analytics screens, while API calls are used to fetch and submit data to the backend. The logic layer is implemented through Flask routes and controller-style functions in app.py, which handle authentication, data retrieval, chatbot processing, and administrator actions. The data layer uses SQLAlchemy models and a SQLite database to persist student accounts, FAQ entries, course data, faculty data, fees, announcements, and chat messages.',
            'The layered design is beneficial because it separates concerns and improves maintainability. Frontend code interacts with backend endpoints through JSON payloads rather than embedding database logic directly in the client. The backend performs validation, enforces role-based access, executes the retrieval pipeline, and maintains integrity across the database. This separation ensures that the system remains understandable and extensible. It also allows future upgrades such as replacing SQLite with PostgreSQL, adding an external authentication provider, or integrating a custom AI service without redesigning the entire user interface.',
            'The system design also includes a clear flow for user interaction. A student logs in through the frontend, the backend verifies the credentials, a session is created, and route guards protect protected resources. When a question is asked, the backend retrieves the FAQ corpus, runs the TF-IDF matching process, selects the best answer, and returns the result as structured JSON. The answer is then displayed in the chat UI. This architecture is simple to understand yet robust enough to deliver a full-featured educational application.'
        ]
    )

    add_section(
        doc,
        '7. Database Design',
        [
            'The database design for SmartCampus is based on SQLite, which is sufficient for a local academic project while still following relational database principles. The schema is organized around a central student table and several supporting information tables. The student table stores name, email, password hash, role, and creation timestamp. This table is linked to chat messages through a one-to-many relationship, allowing administrators to review per-student interactions and enabling the system to preserve a personalized conversation history.',
            'The knowledge tables include faq, course, department, faculty, fee, and announcement. Each table stores relevant operational fields such as question and answer text, description, department names, contact details, fee values, and published dates. These records are treated as structured campus artifacts rather than free-form text, which helps the retrieval system remain transparent and stable. The database structure supports both student-facing use cases and administrative maintenance procedures. It also makes the project easier to extend, because additional records such as calendars, timetables, or libraries can be added without reworking the entire system.',
            'The administrator is responsible for maintaining these tables, while students can only query them through the public portal and their personal chat history. This separation aligns with good application design, because sensitive management operations are available only to the privileged role. The database also saves relevant metadata such as category and confidence score for each chatbot interaction, which supports later analytics and model performance review.'
        ],
        [
            'student: stores user identity and role information.',
            'faq: contains canonical Q&A pairs and aliases for retrieval.',
            'course: stores academic programs and duration information.',
            'department: stores department details and leadership information.',
            'faculty: stores faculty specialization and contact information.',
            'fee: stores academic fee schedules and notes.',
            'announcement: stores campus notices and event updates.',
            'chat_message: stores student questions, generated answers, and metadata.'
        ]
    )

    add_section(
        doc,
        '8. NLP and Chatbot Design',
        [
            'The core intelligence of SmartCampus lies in the NLP-based retrieval pipeline implemented in chatbot.py. The system does not use a hardcoded sequence of if-else rules for every possible query. Instead, it collects a corpus of FAQ entries, builds a searchable representation of question text, and computes the semantic closeness between the user’s question and each stored FAQ entry. This method is known as information retrieval using TF-IDF vectorization and cosine similarity, which is appropriate for a project with modest knowledge size and a requirement for explainable answers.',
            'The first step in the pipeline is text preprocessing. The raw user input and FAQ text are normalized by converting to lowercase and removing punctuation while keeping meaningful alphanumeric content. This reduces noise and prevents minor formatting differences from preventing matches. The system also uses aliases, which allow administrators to add alternate phrasings of the same question so that students can ask in more natural language. The FAQ collection includes category labels such as courses, admissions, fees, faculty, exam schedules, and contact information, which helps organize answers and improve retrieval quality.',
            'In the next stage, scikit-learn’s TfidfVectorizer is used to transform the corpus into numerical vectors. The vectorizer assigns greater importance to terms that are informative within the corpus and less importance to repeated, generic words. The user query is transformed in the same vector space, and cosine similarity is computed between the query vector and each FAQ vector. The answer associated with the highest-scoring match is selected. To avoid false positives, the system checks a minimum confidence threshold. If the best match falls below this threshold, the user receives a fallback response stating that the query could not be matched confidently. This is a critical design feature because it keeps the chatbot reliable and honest.'
        ],
        [
            'Text normalization reduces variations caused by punctuation and capitalization.',
            'TF-IDF highlights important terms in the corpus and ignores irrelevant noise.',
            'Cosine similarity measures the closeness of the user query to each FAQ entry.',
            'Fallback logic ensures the chatbot does not answer with low-confidence or fabricated information.',
            'Question history is stored to support analysis and future improvements.'
        ]
    )

    add_section(
        doc,
        '9. Implementation Details',
        [
            'The implementation of SmartCampus is organized into a small but clearly separated modular structure. The file app.py serves as the application entry point and houses the Flask app factory, route definitions, authentication functions, API endpoints, and session management logic. The backend exposes routes such as /api/auth/login, /api/auth/register, /api/chat, /api/bootstrap, and /api/admin/stats, each responsible for a specific domain of functionality. These routes allow the frontend to interact with the system using JSON-based requests and responses, making the architecture clean and efficient.',
            'The file models.py defines the SQLAlchemy models for Student, FAQ, Course, Department, Faculty, Fee, Announcement, and ChatMessage. Each model includes a to_dict method that serializes the record for API responses. Secure password storage is handled through password hashing utilities so that plaintext passwords are never stored in the database. This is a major security consideration because student credentials are highly sensitive. The login_required decorator and role checks enforce authorization for protected route access. The system identifies students and administrators distinctly through a role field, ensuring that admin-only functionalities remain protected from unauthorized access.',
            'The file seed.py is used to populate the database with realistic demo data. It contains course records, faculty details, department information, admission notices, fee values, and multiple FAQs that reflect typical college queries. The app uses this data when the database is first initialized, which allows the user to interact with a realistic dataset immediately. The file static/js/app.js manages the client-side logic for login, registration, chat rendering, and admin actions. The styling in static/css/styles.css creates a polished interface, while static/index.html hosts the main app shell. This modular division makes the project easy to understand, debug, and extend.'
        ]
    )

    add_section(
        doc,
        '10. User Interface and Experience',
        [
            'The user interface of SmartCampus is designed to be intuitive, modern, and responsive. The landing page presents a clean authentication screen with a student and admin demo access panel. Once logged in, students see a dashboard that includes campus highlights, recent announcements, and a chat box for asking questions. This layout is user-friendly because it groups the most useful functions into one interface without overwhelming the user. The design focuses on readability, minimal clutter, and quick access to information. The responsive styling ensures that the experience remains effective on different screen sizes, including laptops, tablets, and mobile devices.',
            'The chat interface is the centerpiece of the application and is designed to encourage natural interaction. Students enter questions in plain language, and the result is displayed as a message bubble with the answer and the relevant category. Recent questions are retained and can be revisited later, providing continuity in the student experience. The admin interface is distinct from the student experience and includes a management workspace for statistics, quick insights, and CRUD forms. This separation improves clarity and ensures that sensitive admin operations are not mixed into the student layout.',
            'Experience design also includes feedback and usability considerations. When a query is unresolved or too uncertain, the interface does not pretend to know the answer; instead, it returns a fallback response and encourages the user to rephrase the question or contact a department. This honesty improves trust. The combination of a vibrant layout, simple navigation, and role-based workflow makes the system approachable while still demonstrating serious technical depth.'
        ]
    )

    add_section(
        doc,
        '11. Testing, Validation, and Evaluation',
        [
            'Testing is a critical part of the SmartCampus project because the application combines backend logic, database transactions, UI interactions, and NLP matching. The project includes automated regression tests for the chatbot logic and the authentication flow. These tests validate that the retrieval engine matches representative student questions, handles fallback conditions correctly, and respects student-specific history boundaries. For example, the system was tested with questions such as admission procedures, fee inquiries, course listings, faculty contact information, placement assistance, and hostel availability. The test suite confirms that the bot responds appropriately and that session-based access is enforced for protected routes.',
            'The evaluation process also focuses on correctness and confidence. Because retrieval-based systems can produce incorrect matches when the user asks an unrelated or ambiguous question, the project deliberately imposes a threshold for similarity. The fallback response is triggered when the best score is below the acceptable level. This ensures that the system remains cautious rather than overconfident. The same philosophy is used in admin endpoints, where invalid or unauthorized access is rejected. In addition, the project underwent smoke tests to confirm that the Flask app starts correctly, the route endpoints return expected data, and the JavaScript syntax is valid.',
            'The validation results were successful. The tests passed after the database was reset and the app was run with a secure development secret. The Flask application served on the local development URL and the browser was able to render the login screen and access the dashboard. This confirms that the system works end to end and that the essential project requirements are met. This testing stage is important because it shows the project is not only conceptually valid but also practically usable.'
        ]
    )

    add_section(
        doc,
        '12. Security, Challenges, and Limitations',
        [
            'Security is a vital part of any application that stores user credentials and institutional information. SmartCampus addresses this by using hashed passwords instead of storing plain text, by preserving user sessions in secure cookies, and by restricting admin routes to admin-role users only. The application configuration includes a secret key environment variable, which is recommended for deployment settings. This approach prevents accidental exposure of sensitive application data and aligns with standard web application security practices. The design also follows the principle of keeping user-facing operations separate from admin-managed content, limiting the surface area for misuse.',
            'The project also faced several implementation challenges. One challenge was making the chatbot intelligent enough to handle natural-language questions without becoming brittle. To address this, the FAQ corpus was enriched with aliases and category data, and the retrieval threshold was tuned to balance relevance with caution. Another challenge involved keeping the admin dashboard and student experience synchronized with the database while avoiding duplicate or inconsistent records. Data validation in backend routes and schema consistency were critical to handling this effectively. The project also needed to be organized modularly so that future developers could extend the system without confusion.',
            'Although the project is successful for a local educational deployment, it has some limitations. The chatbot is based on a retrieval model rather than a generative AI model, which means it relies on the quality of the curated FAQ dataset. If a question is outside the knowledge base, the fallback response is used. The application also uses SQLite, which is suitable for local development but might not meet the demands of a large-scale production environment. For future scaling, a production-grade database, a stronger authentication system, and a more advanced NLP engine could be explored.'
        ]
    )

    add_section(
        doc,
        '13. Future Scope and Enhancements',
        [
            'SmartCampus is a strong foundation for further development in educational technology. One obvious enhancement is integration with a production-grade database such as PostgreSQL or MySQL to support larger datasets and concurrent users. The application could also be expanded with real-time communication, a mobile application, or a more advanced login system using OAuth or institutional identity providers. Another important enhancement would be multilingual support, allowing students to ask questions in multiple languages and receive culturally relevant responses. This would broaden accessibility for diverse student communities.',
            'A natural next step would be to incorporate a more advanced language model or hybrid retrieval architecture, combining vector search with semantic ranking and contextual responses. However, for a student project, the current TF-IDF model is appropriate because it is transparent, lightweight, and easy to explain. The system could also be extended with a recommendation engine for course suggestions, timetable management, or assignment reminders. In addition, the admin dashboard can be enhanced with charts, alerting, and exportable reports so administrators can make more informed decisions about student support and communication.',
            'The overall future direction is to transform SmartCampus from a demonstrative project into an institutional platform that supports real student engagement. With thoughtful expansion, the system could become a central digital support assistant for campus admissions, academic advising, faculty communication, and general information management. This would improve the communication loop between students and the institution while also demonstrating how AI can be applied responsibly in education.'
        ]
    )

    add_section(
        doc,
        '14. Conclusion',
        [
            'SmartCampus has successfully demonstrated how a college information system can be built using practical open-source technologies and a retrieval-based AI chatbot. The system integrates web development, database management, secure authentication, admin management, and intelligent response generation into a single application that is useful to both students and administrators. It addresses a genuine problem in academic institutions: the need for accessible, accurate, and centralized information. By creating a conversational interface around a structured FAQ knowledge base, the project makes campus information easier to discover and more engaging to use.',
            'The project also highlights the importance of balancing AI functionality with reliability. The retrieval pipeline is not designed to fabricate answers; instead, it uses similarity metrics, category mapping, and a confidence threshold to determine whether an answer can be trusted. This design choice makes the system more suitable for an educational environment where clarity and correctness are essential. The admin features ensure that knowledge remains current and maintainable, while the student interface provides a modern and approachable experience.',
            'In conclusion, SmartCampus represents a realistic and functional academic project with a strong blend of technical depth and practical impact. It meets the goals of a modern college information system by improving communication, simplifying access to information, and illustrating the role of AI in higher education. The project stands as a successful example of how intelligent systems can be designed, developed, and evaluated in a way that is both technically sound and academically valuable.'
        ]
    )

    doc.save('SmartCampus_Project_Report.docx')
    print('Created SmartCampus_Project_Report.docx')


if __name__ == '__main__':
    build_report()
